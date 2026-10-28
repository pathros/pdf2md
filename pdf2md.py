#!/usr/bin/env python3
"""pdf2md - Transcribe PDF o imágenes a Markdown o Word con IA local.

Modos:
  pdf     Extrae texto/tablas con pymupdf4llm. Páginas escaneadas → OCR con
          Ollama (visión). Sin límite de páginas (una a la vez).
  images  Transcribe N imágenes con el mismo OCR y las junta en un archivo.
  word    Convierte .docx → Markdown (pandoc; sin OCR).
  excel   Convierte .xlsx / .xlsm → tablas Markdown (sin OCR).
  pptx    Convierte .pptx / .pptm → Markdown (sin OCR).
  md      Convierte .md → Word (.docx), sin OCR ni pandoc.

Salida: Windows → .docx por defecto; macOS/Linux → .md o .docx (--format).
Excel, PowerPoint y Word salen siempre a .md. Un .md sale siempre a .docx.
"""

from __future__ import annotations

import argparse
import base64
import json
import math
import os
import re
import shlex
import shutil
import subprocess
import sys
import textwrap
from datetime import date, datetime, time
from pathlib import Path
from zipfile import BadZipFile

import requests
from docx import Document
from docx.shared import Inches, Pt
from rich.console import Console, Group
from rich.live import Live
from rich.panel import Panel
from rich.progress import (
    BarColumn,
    Progress,
    SpinnerColumn,
    TextColumn,
    TimeElapsedColumn,
    TimeRemainingColumn,
)
from rich.prompt import Prompt
from rich.text import Text

import pymupdf  # PyMuPDF
import pymupdf4llm
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.exc import PackageNotFoundError

console = Console()

# --- Configuración por defecto (sobreescribible por CLI o variables de entorno) ---
DEFAULT_OLLAMA_HOST = os.environ.get("OLLAMA_HOST", "http://localhost:11434")
# Fallback si Ollama no responde o no hay modelo conocido instalado.
DEFAULT_VISION_MODEL = os.environ.get("PDF2MD_VISION_MODEL", "qwen3.5:9b")

# Preferencia cuando hay varios modelos de visión en Ollama.
# En máquinas con <24 GB de RAM se prioriza el 9b (cabe en 16 GB).
VISION_MODELS_LOW_RAM = (
    "qwen3.5:9b",
    "qwen3.5:latest",
    "qwen3.5:4b",
    "qwen3.6:9b",
    "benhaotang/Nanonets-OCR-s:F16",
)
VISION_MODELS_HIGH_RAM = (
    "qwen3.8:27b",
    "qwen3.6:27b",
    "qwen3.5:27b",
    *VISION_MODELS_LOW_RAM,
)
# Por debajo de esto, un 27B Q4 (~17 GB) pelea con el SO y suele ir a swap.
HIGH_RAM_GB = 24.0

# Si una página tiene menos de estos caracteres de texto real y contiene imágenes,
# se considera escaneada (candidata a OCR).
OCR_TEXT_THRESHOLD = 20

# DPI de render para el OCR con Ollama (zoom 2.5 ~= 180 dpi).
OCR_RENDER_ZOOM = 2.5

IMAGE_EXTENSIONS = frozenset({
    ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff",
})

EXCEL_EXTENSIONS = frozenset({".xlsx", ".xlsm"})

PPTX_EXTENSIONS = frozenset({".pptx", ".pptm", ".ppsx", ".ppsm"})

MD_EXTENSIONS = frozenset({".md", ".markdown", ".mdown"})

_TITLE_PLACEHOLDERS = {
    PP_PLACEHOLDER.TITLE,
    PP_PLACEHOLDER.CENTER_TITLE,
}
if hasattr(PP_PLACEHOLDER, "VERTICAL_TITLE"):
    _TITLE_PLACEHOLDERS.add(PP_PLACEHOLDER.VERTICAL_TITLE)

# Ancho máximo de texto dentro de una celda Markdown. El resto se parte con <br>.
# 0 (por CLI) = no partir. Un salto de línea real rompería la fila de la tabla.
DEFAULT_CELL_WIDTH = 60

# Más de estas filas o columnas vacías seguidas separan otra tabla.
_EXCEL_EMPTY_ROW_GAP = 2
_EXCEL_EMPTY_COL_GAP = 2

OCR_PROMPT = (
    "Eres un motor de OCR. Transcribe TODO el texto visible de esta imagen a "
    "Markdown limpio y bien estructurado. Respeta el idioma original (no traduzcas). "
    "Copia con fidelidad nombres, fechas, cifras, sellos y puntuación. "
    "Usa encabezados (#), listas y tablas de Markdown cuando corresponda y conserva "
    "el orden de lectura. Si un fragmento es ilegible, escribe [ilegible]; no inventes. "
    "NO resumas y NO agregues comentarios ni explicaciones: "
    "devuelve únicamente la transcripción en Markdown."
)

PREVIEW_LINE_COUNT = 5
PREVIEW_LINE_WIDTH = 110

PROGRESS_COLUMNS = (
    SpinnerColumn(),
    TextColumn("[progress.description]{task.description}"),
    BarColumn(),
    TextColumn("{task.completed}/{task.total}"),
    TimeElapsedColumn(),
    TimeRemainingColumn(),
)


def default_output_dir() -> Path:
    env = os.environ.get("PDF2MD_OUTPUT_DIR")
    if env:
        return Path(os.path.expanduser(env))
    if os.name == "nt":
        return Path.home() / "Documents" / "Transcripciones-PDF"
    return Path.home() / "Downloads"


def default_output_format() -> str:
    env = (os.environ.get("PDF2MD_FORMAT") or "").strip().lower()
    if env in ("md", "markdown"):
        return "md"
    if env in ("docx", "doc", "word"):
        return "docx"
    return "docx" if os.name == "nt" else "md"


DEFAULT_OUTPUT_DIR = default_output_dir()


# --------------------------------------------------------------------------- #
# Hardware / modelo de visión
# --------------------------------------------------------------------------- #
def total_ram_gb() -> float | None:
    """RAM física total en GB, o None si no se pudo leer."""
    try:
        if sys.platform == "darwin":
            import subprocess

            out = subprocess.check_output(["sysctl", "-n", "hw.memsize"], text=True)
            return int(out.strip()) / (1024 ** 3)
        if sys.platform == "win32":
            import ctypes

            class MEMORYSTATUSEX(ctypes.Structure):
                _fields_ = [
                    ("dwLength", ctypes.c_ulong),
                    ("dwMemoryLoad", ctypes.c_ulong),
                    ("ullTotalPhys", ctypes.c_ulonglong),
                    ("ullAvailPhys", ctypes.c_ulonglong),
                    ("ullTotalPageFile", ctypes.c_ulonglong),
                    ("ullAvailPageFile", ctypes.c_ulonglong),
                    ("ullTotalVirtual", ctypes.c_ulonglong),
                    ("ullAvailVirtual", ctypes.c_ulonglong),
                    ("ullAvailExtendedVirtual", ctypes.c_ulonglong),
                ]

            stat = MEMORYSTATUSEX()
            stat.dwLength = ctypes.sizeof(MEMORYSTATUSEX)
            if not ctypes.windll.kernel32.GlobalMemoryStatusEx(ctypes.byref(stat)):
                return None
            return stat.ullTotalPhys / (1024 ** 3)
        with open("/proc/meminfo", encoding="utf-8") as fh:
            for line in fh:
                if line.startswith("MemTotal:"):
                    return int(line.split()[1]) / (1024 ** 2)
    except Exception:  # noqa: BLE001
        return None
    return None


def list_ollama_model_names(host: str) -> set[str]:
    try:
        resp = requests.get(f"{host}/api/tags", timeout=5)
        resp.raise_for_status()
    except requests.RequestException:
        return set()
    return {m.get("name", "") for m in resp.json().get("models", []) if m.get("name")}


def model_is_installed(names: set[str], model: str) -> bool:
    """True solo si el tag pedido está instalado (no basta la misma familia)."""
    if model in names:
        return True
    if ":" not in model and f"{model}:latest" in names:
        return True
    return False


def resolve_vision_model(host: str, requested: str | None) -> str:
    """Elige el modelo de visión: CLI / env / lo instalado según RAM."""
    if requested:
        return requested
    env = os.environ.get("PDF2MD_VISION_MODEL")
    if env:
        return env

    ram = total_ram_gb()
    preferred = VISION_MODELS_HIGH_RAM if (ram or 0) >= HIGH_RAM_GB else VISION_MODELS_LOW_RAM
    names = list_ollama_model_names(host)
    for candidate in preferred:
        if model_is_installed(names, candidate):
            return candidate
    return DEFAULT_VISION_MODEL


def warn_if_model_too_heavy(model: str) -> None:
    ram = total_ram_gb()
    if ram is None:
        return
    heavy = any(tag in model.lower() for tag in (":27b", ":35b", ":122b"))
    if heavy and ram < HIGH_RAM_GB:
        console.print(
            f"[yellow]Aviso:[/yellow] '{model}' pesa ~17 GB+ (Q4). "
            f"Esta máquina tiene ~{ram:.0f} GB de RAM; Windows + Ollama "
            "suelen quedarse cortos y el 27b irá a swap o no cargará. "
            "Recomendado: [bold]qwen3.5:9b[/bold] (~6.6 GB)."
        )


# --------------------------------------------------------------------------- #
# Utilidades de ruta / entrada
# --------------------------------------------------------------------------- #
def _strip_wrapping_quotes(value: str) -> str:
    s = value.strip()
    if len(s) >= 2 and s[0] == s[-1] and s[0] in ("'", '"'):
        return s[1:-1]
    return s


def clean_dropped_path(raw: str) -> str:
    """Normaliza una ruta arrastrada a la terminal (Finder, Explorer o PowerShell)."""
    s = raw.strip()
    if s.startswith("&"):
        s = s[1:].strip()
    s = _strip_wrapping_quotes(s)
    if os.name != "nt":
        for esc in ("\\ ", "\\(", "\\)", "\\&", "\\'", "\\\""):
            s = s.replace(esc, esc[1])
    return os.path.expanduser(s.strip())


def parse_dropped_paths(raw: str) -> list[Path]:
    """Separa varias rutas arrastradas (Finder en macOS o Explorer en Windows)."""
    raw = raw.strip()
    if not raw:
        return []
    posix = os.name != "nt"
    to_split = raw.replace("& ", " ") if os.name == "nt" else raw
    try:
        tokens = shlex.split(to_split, posix=posix)
    except ValueError:
        tokens = to_split.split()
    paths: list[Path] = []
    for token in tokens:
        cleaned = clean_dropped_path(token)
        if cleaned:
            paths.append(Path(cleaned))
    return paths


def is_image_path(path: Path) -> bool:
    return path.suffix.lower() in IMAGE_EXTENSIONS


def is_excel_path(path: Path) -> bool:
    return path.suffix.lower() in EXCEL_EXTENSIONS


def is_pptx_path(path: Path) -> bool:
    return path.suffix.lower() in PPTX_EXTENSIONS


def is_md_path(path: Path) -> bool:
    return path.suffix.lower() in MD_EXTENSIONS


def prompt_for_mode() -> str | None:
    """Devuelve el comando elegido, o None si el usuario cancela."""
    console.print(
        Panel.fit(
            "[bold]1[/bold]  PDF → Markdown / Word\n"
            "[bold]2[/bold]  Imágenes → Markdown / Word (OCR con Ollama)\n"
            "[bold]3[/bold]  Word (.docx) → Markdown\n"
            "[bold]4[/bold]  Excel (.xlsx) → Markdown\n"
            "[bold]5[/bold]  PowerPoint (.pptx) → Markdown\n"
            "[bold]6[/bold]  Markdown (.md) → Word\n"
            "[bold]0[/bold]  Cancelar",
            title="pdf2md",
            border_style="cyan",
        )
    )
    while True:
        choice = Prompt.ask(
            "[cyan]Modo[/cyan]",
            choices=[
                "0", "1", "2", "3", "4", "5", "6",
                "pdf", "images", "img", "imgs",
                "word", "docx", "excel", "xlsx", "xlsm",
                "pptx", "ppt", "slides",
                "md", "markdown",
                "cancel", "cancelar", "q",
            ],
            default="1",
            show_choices=False,
        ).strip().lower()
        if choice in ("0", "cancel", "cancelar", "q"):
            return None
        if choice in ("1", "pdf"):
            return "pdf"
        if choice in ("2", "images", "img", "imgs"):
            return "images"
        if choice in ("3", "word", "docx"):
            return "word"
        if choice in ("4", "excel", "xlsx", "xlsm"):
            return "excel"
        if choice in ("5", "pptx", "ppt", "slides"):
            return "pptx"
        if choice in ("6", "md", "markdown"):
            return "md"
        console.print("[yellow]Elige 1, 2, 3, 4, 5, 6 o 0 (Cancelar).[/yellow]")


def prompt_for_pdf() -> Path:
    """Pide interactivamente la ruta del PDF hasta obtener una válida."""
    console.print(
        Panel.fit(
            "[bold]Arrastra el archivo PDF a esta terminal[/bold] y pulsa Enter.\n"
            "También puedes pegar o escribir la ruta.",
            title="pdf2md · PDF",
            border_style="cyan",
        )
    )
    while True:
        raw = Prompt.ask("[cyan]Ruta del PDF[/cyan]")
        if not raw.strip():
            console.print("[yellow]No escribiste nada. Intenta de nuevo.[/yellow]")
            continue
        path = Path(clean_dropped_path(raw))
        if not path.exists():
            console.print(f"[red]No existe:[/red] {path}")
            continue
        if path.suffix.lower() != ".pdf":
            console.print("[red]El archivo no parece un .pdf[/red]")
            continue
        return path


def prompt_for_images() -> list[Path]:
    """Pide una o más rutas de imagen (arrastrar desde Finder o Explorer)."""
    console.print(
        Panel.fit(
            "[bold]Arrastra una o más imágenes a esta terminal[/bold] y pulsa Enter.\n"
            "Las rutas quedan separadas por espacio (Finder/Explorer lo hace solo).\n"
            f"Formatos: {', '.join(sorted(IMAGE_EXTENSIONS))}",
            title="pdf2md · Imágenes",
            border_style="cyan",
        )
    )
    while True:
        raw = Prompt.ask("[cyan]Rutas de imágenes[/cyan]")
        paths = parse_dropped_paths(raw)
        if not paths:
            console.print("[yellow]No escribiste nada. Intenta de nuevo.[/yellow]")
            continue
        ok, err = validate_image_paths(paths)
        if err:
            for msg in err:
                console.print(f"[red]{msg}[/red]")
            continue
        return ok


def prompt_for_docx() -> Path:
    """Pide interactivamente la ruta del .docx hasta obtener una válida."""
    console.print(
        Panel.fit(
            "[bold]Arrastra el archivo Word (.docx) a esta terminal[/bold] y pulsa Enter.\n"
            "También puedes pegar o escribir la ruta.",
            title="pdf2md · Word → Markdown",
            border_style="cyan",
        )
    )
    while True:
        raw = Prompt.ask("[cyan]Ruta del .docx[/cyan]")
        if not raw.strip():
            console.print("[yellow]No escribiste nada. Intenta de nuevo.[/yellow]")
            continue
        path = Path(clean_dropped_path(raw))
        if not path.exists():
            console.print(f"[red]No existe:[/red] {path}")
            continue
        if path.suffix.lower() != ".docx":
            console.print("[red]El archivo no parece un .docx[/red]")
            continue
        return path


def prompt_for_xlsx() -> Path:
    """Pide interactivamente la ruta del .xlsx hasta obtener una válida."""
    console.print(
        Panel.fit(
            "[bold]Arrastra el archivo Excel (.xlsx) a esta terminal[/bold] y pulsa Enter.\n"
            "También puedes pegar o escribir la ruta. Se acepta .xlsm.",
            title="pdf2md · Excel → Markdown",
            border_style="cyan",
        )
    )
    while True:
        raw = Prompt.ask("[cyan]Ruta del .xlsx[/cyan]")
        if not raw.strip():
            console.print("[yellow]No escribiste nada. Intenta de nuevo.[/yellow]")
            continue
        path = Path(clean_dropped_path(raw))
        if not path.exists():
            console.print(f"[red]No existe:[/red] {path}")
            continue
        if path.suffix.lower() == ".xls":
            console.print(
                "[red]El formato .xls (Excel 97-2003) no está soportado.[/red] "
                "Ábrelo y guárdalo como .xlsx."
            )
            continue
        if not is_excel_path(path):
            console.print("[red]El archivo no parece un .xlsx[/red]")
            continue
        return path


def prompt_for_pptx() -> Path:
    """Pide interactivamente la ruta del .pptx hasta obtener una válida."""
    console.print(
        Panel.fit(
            "[bold]Arrastra el archivo PowerPoint (.pptx) a esta terminal[/bold] y pulsa Enter.\n"
            "También puedes pegar o escribir la ruta. Se acepta .pptm.",
            title="pdf2md · PowerPoint → Markdown",
            border_style="cyan",
        )
    )
    while True:
        raw = Prompt.ask("[cyan]Ruta del .pptx[/cyan]")
        if not raw.strip():
            console.print("[yellow]No escribiste nada. Intenta de nuevo.[/yellow]")
            continue
        path = Path(clean_dropped_path(raw))
        if not path.exists():
            console.print(f"[red]No existe:[/red] {path}")
            continue
        if path.suffix.lower() == ".ppt":
            console.print(
                "[red]El formato .ppt (PowerPoint 97-2003) no está soportado.[/red] "
                "Ábrelo y guárdalo como .pptx."
            )
            continue
        if not is_pptx_path(path):
            console.print("[red]El archivo no parece un .pptx[/red]")
            continue
        return path


def prompt_for_md() -> Path:
    """Pide interactivamente la ruta del .md hasta obtener una válida."""
    console.print(
        Panel.fit(
            "[bold]Arrastra el archivo Markdown (.md) a esta terminal[/bold] y pulsa Enter.\n"
            "También puedes pegar o escribir la ruta. Se acepta .markdown.",
            title="pdf2md · Markdown → Word",
            border_style="cyan",
        )
    )
    while True:
        raw = Prompt.ask("[cyan]Ruta del .md[/cyan]")
        if not raw.strip():
            console.print("[yellow]No escribiste nada. Intenta de nuevo.[/yellow]")
            continue
        path = Path(clean_dropped_path(raw))
        if not path.exists():
            console.print(f"[red]No existe:[/red] {path}")
            continue
        if not is_md_path(path):
            console.print("[red]El archivo no parece un .md[/red]")
            continue
        return path


def validate_image_paths(paths: list[Path]) -> tuple[list[Path], list[str]]:
    ok: list[Path] = []
    err: list[str] = []
    seen: set[Path] = set()
    for path in paths:
        resolved = path.expanduser()
        if resolved in seen:
            continue
        seen.add(resolved)
        if not resolved.exists():
            err.append(f"No existe: {resolved}")
            continue
        if not resolved.is_file():
            err.append(f"No es un archivo: {resolved}")
            continue
        if not is_image_path(resolved):
            err.append(f"No es una imagen soportada: {resolved}")
            continue
        ok.append(resolved)
    return ok, err


# --------------------------------------------------------------------------- #
# Detección de páginas escaneadas
# --------------------------------------------------------------------------- #
def detect_scanned_pages(pdf_path: Path) -> list[bool]:
    """Devuelve, por página, si parece escaneada (sin capa de texto, con imagen)."""
    doc = pymupdf.open(pdf_path)
    flags: list[bool] = []
    try:
        for page in doc:
            text = page.get_text("text").strip()
            has_images = bool(page.get_images(full=True))
            flags.append(len(text) < OCR_TEXT_THRESHOLD and has_images)
    finally:
        doc.close()
    return flags


# --------------------------------------------------------------------------- #
# OCR con Ollama (modelo de visión) + vista en vivo
# --------------------------------------------------------------------------- #
class TranscriptionDisplay:
    """Barra de progreso + últimas N líneas (estilo log de Docker)."""

    def __init__(self, total: int, label: str = "Procesando") -> None:
        self.progress = Progress(*PROGRESS_COLUMNS, expand=True)
        self.task = self.progress.add_task(label, total=total)
        self._buffer = ""
        self._live: Live | None = None

    def _preview_lines(self) -> list[str]:
        lines = self._buffer.splitlines()
        last = lines[-PREVIEW_LINE_COUNT:] if lines else []
        last = [ln[:PREVIEW_LINE_WIDTH] for ln in last]
        while len(last) < PREVIEW_LINE_COUNT:
            last.append("")
        return last

    def _renderable(self) -> Group:
        body = "\n".join(self._preview_lines()) or " "
        panel = Panel(
            Text(body, style="dim"),
            title="Transcripción (en vivo)",
            subtitle="últimas 5 líneas",
            border_style="cyan",
            height=PREVIEW_LINE_COUNT + 2,
        )
        return Group(self.progress, panel)

    def _refresh(self) -> None:
        if self._live is not None:
            self._live.update(self._renderable())

    def __enter__(self) -> TranscriptionDisplay:
        self._live = Live(
            self._renderable(),
            console=console,
            refresh_per_second=8,
            vertical_overflow="crop",
        )
        self._live.__enter__()
        return self

    def __exit__(self, *args: object) -> None:
        if self._live is not None:
            self._live.__exit__(*args)

    def set_status(self, description: str) -> None:
        self.progress.update(self.task, description=description)
        self._refresh()

    def start_unit(self, description: str) -> None:
        self._buffer = ""
        self.set_status(description)

    def feed(self, chunk: str) -> None:
        self._buffer += chunk
        self._refresh()

    def show_text(self, text: str) -> None:
        self._buffer = text.strip()
        self._refresh()

    def advance(self) -> None:
        self.progress.advance(self.task)
        self._refresh()


def render_page_png(page: "pymupdf.Page", zoom: float = OCR_RENDER_ZOOM) -> bytes:
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
    return pix.tobytes("png")


def check_ollama(host: str, model: str) -> bool:
    """Verifica que Ollama responda y que el tag del modelo esté instalado."""
    names = list_ollama_model_names(host)
    if not names:
        return False
    return model_is_installed(names, model)


def ocr_image_bytes_with_ollama(
    image_bytes: bytes,
    model: str,
    host: str,
    on_chunk: object | None = None,
) -> str:
    payload = {
        "model": model,
        "prompt": OCR_PROMPT,
        "images": [base64.b64encode(image_bytes).decode("ascii")],
        "stream": True,
        # Qwen 3.5/3.6 "piensan" por defecto; para OCR queremos solo la transcripción.
        "think": False,
        "options": {"temperature": 0},
    }
    pieces: list[str] = []
    with requests.post(
        f"{host}/api/generate",
        json=payload,
        stream=True,
        timeout=(15, 1800),
    ) as resp:
        resp.raise_for_status()
        for raw_line in resp.iter_lines(decode_unicode=True):
            if not raw_line:
                continue
            try:
                data = json.loads(raw_line)
            except json.JSONDecodeError:
                continue
            token = data.get("response") or ""
            if token:
                pieces.append(token)
                if callable(on_chunk):
                    on_chunk(token)
            if data.get("done"):
                break
    return "".join(pieces).strip()


def ocr_page_with_ollama(
    png_bytes: bytes,
    model: str,
    host: str,
    on_chunk: object | None = None,
) -> str:
    return ocr_image_bytes_with_ollama(png_bytes, model, host, on_chunk=on_chunk)


# --------------------------------------------------------------------------- #
# Markdown → Word (.docx)
# --------------------------------------------------------------------------- #
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
_UL_RE = re.compile(r"^([ \t]*)[-*+]\s+(.*)$")
_OL_RE = re.compile(r"^([ \t]*)(\d+)[.)]\s+(.*)$")
_PAGE_MARK_RE = re.compile(r"^<!--\s*página\s+\d+\s*-->\s*$", re.I)
_HTML_COMMENT_RE = re.compile(r"^<!--.*?-->\s*$")
_TABLE_ROW_RE = re.compile(r"^\|(.+)\|\s*$")
_TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
_BR_RE = re.compile(r"(?i)<br\s*/?>")
_IMAGE_RE = re.compile(r"!\[([^\]]*)\]\([^)]+\)")
_LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


def _strip_frontmatter(md: str) -> str:
    if not md.startswith("---"):
        return md
    end = md.find("\n---\n", 3)
    if end == -1:
        return md
    return md[end + 5 :]


def _restore_md_text(text: str) -> str:
    """Deshace el escape de tablas (Excel/PowerPoint) y los ``<br>``."""
    text = _BR_RE.sub("\n", text)
    text = (
        text.replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&#39;", "'")
    )
    text = text.replace("\\|", "|").replace("\\`", "`").replace("\\*", "*")
    return text.replace("\\\\", "\\")


def _plain_md_text(text: str) -> str:
    """Texto de un encabezado: sin marcas inline ni saltos."""
    text = _restore_md_text(text)
    text = _IMAGE_RE.sub(lambda m: m.group(1) or "[imagen]", text)
    text = _LINK_RE.sub(r"\1 (\2)", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return " ".join(text.split())


def _add_run_with_breaks(paragraph, text: str, *, bold: bool = False, italic: bool = False) -> None:
    pieces = text.split("\n")
    for idx, piece in enumerate(pieces):
        if idx:
            paragraph.add_run().add_break()
        if not piece:
            continue
        run = paragraph.add_run(piece)
        if bold:
            run.bold = True
        if italic:
            run.italic = True


def _add_inline_runs(paragraph, text: str) -> None:
    text = _restore_md_text(text)
    text = _IMAGE_RE.sub(lambda m: m.group(1) or "[imagen]", text)
    text = _LINK_RE.sub(r"\1 (\2)", text)
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) >= 4:
            _add_run_with_breaks(paragraph, part[2:-2], bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) >= 2:
            _add_run_with_breaks(paragraph, part[1:-1], italic=True)
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            _add_run_with_breaks(paragraph, part[1:-1], italic=True)
        else:
            _add_run_with_breaks(paragraph, part)


def _list_level(indent: str) -> int:
    return len(indent.replace("\t", "    ")) // 2


def _apply_list_level(paragraph, level: int) -> None:
    if level > 0:
        paragraph.paragraph_format.left_indent = Inches(0.25 + 0.35 * level)


def _split_table_row(line: str) -> list[str]:
    """Parte una fila ``| a | b |``. Un ``\\|`` dentro de la celda no es separador."""
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    cells: list[str] = []
    buf: list[str] = []
    escaped = False
    for ch in inner:
        if escaped:
            buf.append("\\")
            buf.append(ch)
            escaped = False
            continue
        if ch == "\\":
            escaped = True
            continue
        if ch == "|":
            cells.append("".join(buf).strip())
            buf = []
            continue
        buf.append(ch)
    if escaped:
        buf.append("\\")
    cells.append("".join(buf).strip())
    return cells


def _set_cell_text(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    _add_inline_runs(paragraph, text)


def _is_plain_line(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    if stripped.startswith(("```", "~~~", ">")):
        return False
    if stripped in ("---", "***", "___"):
        return False
    if (
        _PAGE_MARK_RE.match(stripped)
        or _HTML_COMMENT_RE.match(stripped)
        or _HEADING_RE.match(stripped)
        or _TABLE_ROW_RE.match(stripped)
        or _TABLE_SEP_RE.match(stripped)
        or _UL_RE.match(line)
        or _OL_RE.match(line)
    ):
        return False
    return True


def markdown_to_docx(md: str, path: Path, *, join_lines: bool = False) -> None:
    """Convierte Markdown a un .docx editable en Word.

    ``join_lines`` une líneas seguidas de un mismo párrafo. El export de un
    PDF lo deja en falso: ahí cada línea ya es un bloque.
    """
    doc = Document()
    body = _strip_frontmatter(md)
    lines = body.splitlines()
    i = 0
    page_breaks = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if _PAGE_MARK_RE.match(stripped):
            if page_breaks:
                doc.add_page_break()
            page_breaks += 1
            i += 1
            continue

        if _HTML_COMMENT_RE.match(stripped):
            i += 1
            continue

        if stripped.startswith(("```", "~~~")):
            fence = stripped[:3]
            i += 1
            code_lines: list[str] = []
            while i < len(lines) and not lines[i].strip().startswith(fence):
                code_lines.append(lines[i])
                i += 1
            if i < len(lines):
                i += 1
            p = doc.add_paragraph()
            run = p.add_run("\n".join(code_lines))
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            continue

        heading = _HEADING_RE.match(stripped)
        if heading:
            level = min(len(heading.group(1)), 9)
            text = _plain_md_text(heading.group(2))
            if level == 1 and page_breaks <= 1 and i < 8:
                doc.add_heading(text, level=0)
            else:
                doc.add_heading(text, level=level)
            i += 1
            continue

        if (
            _TABLE_ROW_RE.match(stripped)
            and i + 1 < len(lines)
            and _TABLE_SEP_RE.match(lines[i + 1].strip())
        ):
            rows: list[list[str]] = [_split_table_row(stripped)]
            i += 2
            while i < len(lines) and _TABLE_ROW_RE.match(lines[i].strip()):
                rows.append(_split_table_row(lines[i].strip()))
                i += 1
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx in range(cols):
                    cell_text = row[c_idx] if c_idx < len(row) else ""
                    _set_cell_text(table.rows[r_idx].cells[c_idx], cell_text)
            continue

        ul = _UL_RE.match(line)
        if ul:
            p = doc.add_paragraph(style="List Bullet")
            _apply_list_level(p, _list_level(ul.group(1)))
            _add_inline_runs(p, ul.group(2))
            i += 1
            continue

        ol = _OL_RE.match(line)
        if ol:
            p = doc.add_paragraph(style="List Number")
            _apply_list_level(p, _list_level(ol.group(1)))
            _add_inline_runs(p, ol.group(3))
            i += 1
            continue

        if stripped in ("---", "***", "___"):
            i += 1
            continue

        if stripped.startswith(">"):
            quote: list[str] = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(re.sub(r"^>\s?", "", lines[i].strip()))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            _add_inline_runs(p, " ".join(quote))
            for run in p.runs:
                run.italic = True
            continue

        parts = [stripped]
        i += 1
        if join_lines:
            while i < len(lines) and _is_plain_line(lines[i]):
                parts.append(lines[i].strip())
                i += 1
        p = doc.add_paragraph()
        _add_inline_runs(p, " ".join(parts))

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


def output_extension(fmt: str) -> str:
    return ".docx" if fmt == "docx" else ".md"


def resolve_output_format(explicit: str | None, output: str | None) -> str:
    if explicit in ("md", "docx"):
        return explicit
    if output:
        suffix = Path(os.path.expanduser(output)).suffix.lower()
        if suffix == ".docx":
            return "docx"
        if suffix == ".md":
            return "md"
    return default_output_format()


def resolve_output_path(source_stem: str, output: str | None, fmt: str) -> Path:
    ext = output_extension(fmt)
    if output:
        out = Path(os.path.expanduser(output))
        looks_like_dir = (
            out.is_dir()
            or output.endswith(("/", "\\"))
            or out.suffix == ""
        )
        if looks_like_dir:
            return out / f"{source_stem}{ext}"
        if out.suffix.lower() not in (".md", ".docx"):
            out = out.with_suffix(ext)
        return out
    dest = default_output_dir()
    dest.mkdir(parents=True, exist_ok=True)
    return dest / f"{source_stem}{ext}"


def write_output(markdown: str, out_path: Path, fmt: str) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx" or out_path.suffix.lower() == ".docx":
        markdown_to_docx(markdown, out_path)
        return
    out_path.write_text(markdown, encoding="utf-8")


# --------------------------------------------------------------------------- #
# Word (.docx) → Markdown
# --------------------------------------------------------------------------- #
def find_pandoc() -> str | None:
    return shutil.which("pandoc")


def docx_to_markdown(docx_path: Path) -> str:
    """Convierte .docx a Markdown con pandoc (mejor fidelidad de estructura)."""
    pandoc = find_pandoc()
    if not pandoc:
        raise RuntimeError(
            "Se necesita 'pandoc' para Word → Markdown. "
            "Instálalo con: brew install pandoc"
        )
    result = subprocess.run(
        [
            pandoc,
            str(docx_path),
            "-f", "docx",
            "-t", "gfm",
            "--wrap=none",
        ],
        check=False,
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        err = (result.stderr or result.stdout or "error desconocido").strip()
        raise RuntimeError(f"pandoc falló: {err}")
    body = result.stdout.strip()
    frontmatter = (
        "---\n"
        f"fuente: {docx_path.name}\n"
        "conversion: word→markdown (pandoc)\n"
        "generado_por: pdf2md\n"
        "---\n\n"
        f"# {docx_path.stem}\n\n"
    )
    return frontmatter + body + "\n"


# --------------------------------------------------------------------------- #
# Excel (.xlsx / .xlsm) → Markdown
# --------------------------------------------------------------------------- #
def _normalize_newlines(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    for sep in ("\u2028", "\u2029", "\v", "\f"):
        text = text.replace(sep, "\n")
    return text


def _strip_control_chars(text: str) -> str:
    """Quita controles que no son salto de línea ni tabulación."""
    kept: list[str] = []
    for ch in text:
        code = ord(ch)
        if ch in "\n\t" or (code >= 32 and code != 127):
            kept.append(ch)
        else:
            kept.append(" ")
    return "".join(kept)


def _format_has_percent(number_format: str | None) -> bool:
    if not number_format:
        return False
    unquoted = re.sub(r'"[^"]*"', "", number_format)
    return "%" in unquoted


def _format_float(value: float) -> str:
    if math.isnan(value) or math.isinf(value):
        return ""
    if abs(value) >= 1e15 or (value != 0 and abs(value) < 1e-9):
        return format(value, ".12g")
    rounded = round(value, 10)
    if abs(rounded - round(rounded)) < 1e-9:
        return str(int(round(rounded)))
    return f"{rounded:.10f}".rstrip("0").rstrip(".")


def _format_percent(value: float) -> str:
    return _format_float(value * 100) + "%"


def _rich_text_to_str(value: object) -> str | None:
    if type(value).__name__ != "CellRichText":
        return None
    parts: list[str] = []
    for block in value:  # type: ignore[attr-defined]
        if isinstance(block, str):
            parts.append(block)
        else:
            parts.append(str(getattr(block, "text", block)))
    return "".join(parts)


def _as_formula_text(value: object) -> str:
    if isinstance(value, str) and value.startswith("="):
        return value
    text = getattr(value, "text", None)
    if isinstance(text, str) and text.startswith("="):
        return text
    return ""


def _excel_value_to_text(value: object, number_format: str | None) -> str:
    """Valor de celda → texto fiel, sin símbolos de tabla Markdown."""
    if value is None:
        return ""
    rich = _rich_text_to_str(value)
    if rich is not None:
        return rich
    formula = _as_formula_text(value)
    if formula:
        return formula
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            return value.isoformat(sep=" ", timespec="seconds")
        if (
            value.hour == 0
            and value.minute == 0
            and value.second == 0
            and value.microsecond == 0
        ):
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        if _format_has_percent(number_format):
            return _format_percent(value)
        return _format_float(value)
    if isinstance(value, int):
        if _format_has_percent(number_format):
            return _format_percent(float(value))
        return str(value)
    return str(value)


def _cell_has_content(cell: object) -> bool:
    if cell is None:
        return False
    link = getattr(cell, "hyperlink", None)
    if link is not None and getattr(link, "target", None):
        return True
    value = getattr(cell, "value", None)
    if value is None:
        return False
    if isinstance(value, str) and not value.strip() and not _as_formula_text(value):
        return False
    return True


def _hyperlink_target(cell: object) -> str:
    if cell is None:
        return ""
    link = getattr(cell, "hyperlink", None)
    if link is None:
        return ""
    target = getattr(link, "target", None) or ""
    return str(target).strip()


def _render_excel_cell(value_cell: object, formula_cell: object) -> str:
    value = getattr(value_cell, "value", None) if value_cell is not None else None
    number_format = ""
    for cell in (value_cell, formula_cell):
        fmt = getattr(cell, "number_format", None) if cell is not None else None
        if fmt and fmt != "General":
            number_format = str(fmt)
            break
        if fmt and not number_format:
            number_format = str(fmt)
    if value is None and formula_cell is not None:
        raw = getattr(formula_cell, "value", None)
        if raw is not None:
            value = _as_formula_text(raw) or raw
    text = _excel_value_to_text(value, number_format or None)
    target = _hyperlink_target(formula_cell) or _hyperlink_target(value_cell)
    if target and target not in text:
        text = f"{text} ({target})" if text else target
    return text


def _one_line(text: str) -> str:
    text = _strip_control_chars(_normalize_newlines(text))
    parts = [part.strip() for part in text.split("\n")]
    return " ".join(part for part in parts if part)


def _paragraph_text(text: str) -> str:
    """Texto suelto (título sobre la tabla): no debe iniciar un encabezado."""
    text = _one_line(text)
    text = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    if text.startswith("#"):
        text = "\\" + text
    return text


def _escape_md_fragment(text: str) -> str:
    """Escapa lo que parte una fila de tabla Markdown.

    El orden importa: primero ``\\``, para que un ``\\`` final no se coma el
    ``|`` de cierre, y para no duplicar las barras que añadimos después.
    """
    text = text.replace("\\", "\\\\")
    text = text.replace("|", "\\|")
    text = text.replace("`", "\\`")
    text = text.replace("*", "\\*")
    return text


def _wrap_line(line: str, width: int) -> list[str]:
    if width <= 0 or len(line) <= width:
        return [line]
    wrapped = textwrap.wrap(
        line,
        width=width,
        break_long_words=True,
        break_on_hyphens=False,
        replace_whitespace=False,
        drop_whitespace=False,
    )
    return wrapped or [line]


def escape_table_cell(text: str, width: int) -> str:
    """Deja una celda segura dentro de una fila ``| … |``.

    Los saltos reales se vuelven ``<br>`` (un salto de línea partiría la fila).
    El texto más ancho que ``width`` se parte igual, en un límite de palabra.
    ``<br>`` se inserta después de escapar el HTML del usuario.
    """
    text = _strip_control_chars(_normalize_newlines(text)).replace("\t", " ")
    if not text.strip():
        return ""
    text = (
        text.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    pieces: list[str] = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            pieces.append("")
            continue
        pieces.extend(part.strip() for part in _wrap_line(line, width))
    while pieces and not pieces[0]:
        pieces.pop(0)
    while pieces and not pieces[-1]:
        pieces.pop()
    return "<br>".join(_escape_md_fragment(part) for part in pieces)


def _rows_to_table(rows: list[list[str]], width: int) -> str:
    ncol = max(len(row) for row in rows)
    rendered: list[list[str]] = []
    for row in rows:
        padded = list(row) + [""] * (ncol - len(row))
        rendered.append([escape_table_cell(cell, width) for cell in padded])
    while ncol > 1 and all(not rendered[i][ncol - 1] for i in range(len(rendered))):
        ncol -= 1
        for row in rendered:
            row.pop()

    def fmt(cells: list[str]) -> str:
        return "| " + " | ".join(cells) + " |"

    lines = [
        fmt(rendered[0]),
        "| " + " | ".join("---" for _ in range(ncol)) + " |",
    ]
    lines.extend(fmt(row) for row in rendered[1:])
    return "\n".join(lines)


def _split_index_runs(values: list[int], max_gap: int) -> list[list[int]]:
    """Parte índices ordenados cuando el hueco vacío supera ``max_gap``."""
    if not values:
        return []
    groups: list[list[int]] = [[values[0]]]
    for prev, cur in zip(values, values[1:]):
        if cur - prev - 1 > max_gap:
            groups.append([cur])
        else:
            groups[-1].append(cur)
    return groups


def _column_groups(used_cols: list[int], max_gap: int) -> list[list[int]]:
    """Agrupa columnas. Un hueco corto se conserva; uno largo abre otra tabla."""
    if not used_cols:
        return []
    groups: list[list[int]] = [[used_cols[0]]]
    for prev, cur in zip(used_cols, used_cols[1:]):
        hole = cur - prev - 1
        if hole > max_gap:
            groups.append([cur])
        elif hole > 0:
            groups[-1].extend(range(prev + 1, cur + 1))
        else:
            groups[-1].append(cur)
    return groups


def _split_preamble(rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """Filas de un solo valor antes del encabezado pasan a párrafos."""
    counts = [sum(1 for cell in row if cell.strip()) for row in rows]
    if not any(counts):
        return [], []
    if sum(1 for n in counts if n) == 1 and max(counts) == 1:
        text = next(cell.strip() for row in rows for cell in row if cell.strip())
        return [text], []
    header_at = next((i for i, n in enumerate(counts) if n >= 2), None)
    if header_at is None:
        return [], rows
    preamble: list[str] = []
    for row in rows[:header_at]:
        texts = [cell.strip() for cell in row if cell.strip()]
        if len(texts) == 1:
            preamble.append(texts[0])
    return preamble, rows[header_at:]


def _tables_from_grid(
    grid: dict[tuple[int, int], str],
) -> list[tuple[list[str], list[list[str]]]]:
    if not grid:
        return []
    blocks: list[tuple[list[str], list[list[str]]]] = []
    for row_group in _split_index_runs(sorted({r for r, _c in grid}), _EXCEL_EMPTY_ROW_GAP):
        row_set = set(row_group)
        cols_used = sorted({c for r, c in grid if r in row_set})
        for col_group in _column_groups(cols_used, _EXCEL_EMPTY_COL_GAP):
            col_index = {c: i for i, c in enumerate(col_group)}
            raw_rows: list[list[str]] = []
            for r in row_group:
                if not any((r, c) in grid for c in col_group):
                    continue
                line = [""] * len(col_group)
                for c in col_group:
                    if (r, c) in grid:
                        line[col_index[c]] = grid[(r, c)]
                raw_rows.append(line)
            if not raw_rows:
                continue
            preamble, body = _split_preamble(raw_rows)
            if preamble or body:
                blocks.append((preamble, body))
    return blocks


def _heading_text(name: str) -> str:
    cleaned = "".join(
        ch if ord(ch) >= 32 and ord(ch) != 127 else " "
        for ch in name
    )
    return " ".join(cleaned.split()) or "Hoja"


def _yaml_quote(value: str) -> str:
    escaped = (
        value.replace("\\", "\\\\")
        .replace('"', '\\"')
        .replace("\n", "\\n")
        .replace("\r", "")
    )
    return f'"{escaped}"'


def _worksheet_cells(ws: object) -> dict:
    # Solo celdas reales. iter_rows() sigue el dimension del XML, que Excel
    # a veces infla hasta la última columna y haría una tabla gigante vacía.
    cells = getattr(ws, "_cells", None)
    return cells if isinstance(cells, dict) else {}


def _fill_merged(ws: object, grid: dict[tuple[int, int], str]) -> None:
    """Repite el valor de un rango combinado en cada celda.

    Un merge de una sola fila y muchas columnas suele ser un título: se deja
    solo en la celda de origen para que quede como párrafo y no ensanche la tabla.
    """
    merged = getattr(ws, "merged_cells", None)
    ranges = getattr(merged, "ranges", None) if merged is not None else None
    if not ranges:
        return
    for cell_range in list(ranges):
        n_rows = cell_range.max_row - cell_range.min_row + 1
        n_cols = cell_range.max_col - cell_range.min_col + 1
        if n_rows * n_cols > 64 or (n_rows == 1 and n_cols > 3):
            continue
        origin = (cell_range.min_row, cell_range.min_col)
        text = grid.get(origin, "")
        if not str(text).strip():
            continue
        for row in range(cell_range.min_row, cell_range.max_row + 1):
            for col in range(cell_range.min_col, cell_range.max_col + 1):
                grid[(row, col)] = text


def _sheet_grid(ws_values: object, ws_formulas: object) -> dict[tuple[int, int], str]:
    value_cells = _worksheet_cells(ws_values)
    formula_cells = _worksheet_cells(ws_formulas)
    coords = {
        key for key, cell in value_cells.items() if _cell_has_content(cell)
    }
    coords.update(
        key for key, cell in formula_cells.items() if _cell_has_content(cell)
    )
    grid: dict[tuple[int, int], str] = {}
    for key in coords:
        text = _render_excel_cell(value_cells.get(key), formula_cells.get(key))
        if text.strip():
            grid[key] = text
    _fill_merged(ws_formulas, grid)
    return grid


def _render_sheet(
    name: str,
    blocks: list[tuple[list[str], list[list[str]]]],
    width: int,
) -> str:
    parts = [f"## {_heading_text(name)}", ""]
    for preamble, rows in blocks:
        for line in preamble:
            paragraph = _paragraph_text(line)
            if paragraph:
                parts.append(paragraph)
                parts.append("")
        if rows:
            parts.append(_rows_to_table(rows, width))
            parts.append("")
    return "\n".join(parts).rstrip() + "\n"


def xlsx_to_markdown(xlsx_path: Path, cell_width: int = DEFAULT_CELL_WIDTH) -> str:
    """Convierte cada hoja visible en una o más tablas Markdown.

    Las fórmulas usan el valor calculado que Excel guardó. Si no hay caché,
    se escribe la fórmula (``=A1+B1``) para no perder la celda.
    """
    if cell_width < 0:
        cell_width = 0
    try:
        wb_values = load_workbook(xlsx_path, data_only=True)
        wb_formulas = load_workbook(xlsx_path, data_only=False)
    except (InvalidFileException, BadZipFile) as exc:
        raise RuntimeError(
            "No se pudo leer el Excel. Si es .xls antiguo, ábrelo y guárdalo "
            f"como .xlsx. Si tiene contraseña, quítala. Detalle: {exc}"
        ) from exc

    sections: list[str] = []
    try:
        names = list(wb_formulas.sheetnames)
        visible = [
            name for name in names
            if wb_formulas[name].sheet_state == "visible"
        ]
        chosen = visible or names
        for name in chosen:
            grid = _sheet_grid(wb_values[name], wb_formulas[name])
            blocks = _tables_from_grid(grid)
            if not blocks:
                console.print(f"  · {name}: sin datos", markup=False)
                continue
            n_tables = sum(1 for _pre, rows in blocks if rows)
            n_rows = sum(len(rows) for _pre, rows in blocks if rows)
            if n_tables:
                detail = f"{n_rows} fila(s), {n_tables} tabla(s)"
            else:
                detail = "solo texto, sin tabla"
            console.print(f"  · {name}: {detail}", markup=False)
            sections.append(_render_sheet(name, blocks, cell_width))
    finally:
        wb_values.close()
        wb_formulas.close()

    if sections:
        body = "\n".join(section.rstrip() for section in sections).strip() + "\n"
    else:
        body = "_El libro no tiene celdas con datos._\n"

    frontmatter = (
        "---\n"
        f"fuente: {_yaml_quote(xlsx_path.name)}\n"
        f"hojas: {len(sections)}\n"
        'conversion: "excel→markdown"\n'
        "generado_por: pdf2md\n"
        "---\n\n"
        f"# {xlsx_path.stem}\n\n"
    )
    return frontmatter + body


# --------------------------------------------------------------------------- #
# PowerPoint (.pptx / .pptm) → Markdown
# --------------------------------------------------------------------------- #
def _shape_key(shape: object) -> tuple[int, int]:
    top = getattr(shape, "top", None)
    left = getattr(shape, "left", None)
    return (int(top or 0), int(left or 0))


def _ordered_shapes(shapes: object) -> list:
    """Orden de lectura: de arriba hacia abajo y, en un grupo, sus hijos."""
    try:
        ordered = sorted(shapes, key=_shape_key)
    except TypeError:
        ordered = list(shapes)
    result: list = []
    for shape in ordered:
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            result.extend(_ordered_shapes(shape.shapes))
        else:
            result.append(shape)
    return result


def _placeholder_kind(shape: object):
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        return shape.placeholder_format.type
    except (AttributeError, ValueError):
        return None


def _inline_text(paragraph: object) -> str:
    runs = list(getattr(paragraph, "runs", []) or [])
    if not runs:
        return str(getattr(paragraph, "text", "") or "")
    chunks: list[str] = []
    for run in runs:
        text = str(getattr(run, "text", "") or "")
        address = ""
        try:
            link = run.hyperlink
            if link is not None and link.address:
                address = str(link.address).strip()
        except (AttributeError, KeyError, ValueError):
            address = ""
        if address and address not in text:
            text = f"{text} ({address})" if text.strip() else address
        chunks.append(text)
    return "".join(chunks)


def _frame_plain(frame: object) -> str:
    parts = [_inline_text(paragraph) for paragraph in frame.paragraphs]
    return _one_line(" ".join(parts))


def _frame_bullets(frame: object) -> str:
    lines: list[str] = []
    for paragraph in frame.paragraphs:
        text = _paragraph_text(_inline_text(paragraph))
        if not text:
            continue
        level = getattr(paragraph, "level", 0) or 0
        try:
            level = int(level)
        except (TypeError, ValueError):
            level = 0
        level = min(max(level, 0), 8)
        lines.append(f"{'  ' * level}- {text}")
    return "\n".join(lines)


def _picture_alt(shape: object) -> str:
    try:
        descr = shape._element._nvXxPr.cNvPr.get("descr") or ""
    except AttributeError:
        return ""
    return str(descr).strip()


def _pptx_table_markdown(shape: object, width: int) -> str:
    table = shape.table
    grid: list[list[str]] = []
    origins: list[tuple[int, int, int, int, str]] = []
    for r_idx, row in enumerate(table.rows):
        line: list[str] = []
        for c_idx, cell in enumerate(row.cells):
            line.append(cell.text or "")
            if getattr(cell, "is_merge_origin", False):
                span_h = int(getattr(cell, "span_height", 1) or 1)
                span_w = int(getattr(cell, "span_width", 1) or 1)
                origins.append((r_idx, c_idx, span_h, span_w, cell.text or ""))
        grid.append(line)
    for r_idx, c_idx, span_h, span_w, text in origins:
        if not text.strip():
            continue
        if span_h * span_w > 64 or (span_h == 1 and span_w > 3):
            continue
        for rr in range(r_idx, min(r_idx + span_h, len(grid))):
            for cc in range(c_idx, min(c_idx + span_w, len(grid[rr]))):
                if not grid[rr][cc].strip():
                    grid[rr][cc] = text
    rows = [row for row in grid if any(cell.strip() for cell in row)]
    if not rows:
        return ""
    return _rows_to_table(rows, width)


def _chart_table(chart: object, width: int) -> str:
    plots = list(chart.plots)
    if not plots:
        return ""
    plot = plots[0]
    categories = [_one_line(str(category)) for category in plot.categories]
    series_rows: list[tuple[str, list[str]]] = []
    for series in plot.series:
        raw_name = series.name
        if isinstance(raw_name, str):
            name = _one_line(raw_name)
        elif raw_name is None:
            name = ""
        else:
            name = _one_line(str(raw_name))
        values: list[str] = []
        for value in series.values:
            if value is None:
                values.append("")
            else:
                values.append(_excel_value_to_text(value, None))
        series_rows.append((name, values))
    if not categories or not series_rows:
        return ""
    rows = [["", *[name for name, _values in series_rows]]]
    for index, category in enumerate(categories):
        row = [category]
        for _name, values in series_rows:
            row.append(values[index] if index < len(values) else "")
        rows.append(row)
    if not any(cell.strip() for row in rows for cell in row):
        return ""
    return _rows_to_table(rows, width)


def _chart_to_markdown(shape: object, width: int) -> str:
    try:
        chart = shape.chart
    except (AttributeError, ValueError):
        return ""
    title = ""
    try:
        if chart.has_title:
            title = _one_line(chart.chart_title.text_frame.text)
    except Exception:  # noqa: BLE001
        title = ""
    table = ""
    try:
        table = _chart_table(chart, width)
    except Exception:  # noqa: BLE001
        table = ""
    parts: list[str] = []
    if title:
        parts.append(_paragraph_text(title))
    if table:
        parts.append(table)
    return "\n\n".join(parts)


def _shape_markdown(shape: object, width: int) -> str:
    if getattr(shape, "has_table", False):
        return _pptx_table_markdown(shape, width)
    if getattr(shape, "has_chart", False):
        return _chart_to_markdown(shape, width)
    if getattr(shape, "has_text_frame", False):
        return _frame_bullets(shape.text_frame)
    alt = _picture_alt(shape)
    if alt:
        return _paragraph_text(alt)
    return ""


def _notes_markdown(slide: object) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    lines = [
        _paragraph_text(_inline_text(paragraph))
        for paragraph in frame.paragraphs
    ]
    lines = [line for line in lines if line]
    if not lines:
        return ""
    return "### Notas\n\n" + "\n\n".join(lines)


def _slide_is_hidden(prs: object, index: int) -> bool:
    try:
        slide_id = prs.slides._sldIdLst[index]
    except (AttributeError, IndexError):
        return False
    return slide_id.get("show") == "0"


def _render_slide(slide: object, number: int, hidden: bool, width: int) -> str:
    title = ""
    subtitle = ""
    blocks: list[str] = []
    for shape in _ordered_shapes(slide.shapes):
        kind = _placeholder_kind(shape)
        if (
            kind in _TITLE_PLACEHOLDERS
            and getattr(shape, "has_text_frame", False)
            and not title
        ):
            title = _frame_plain(shape.text_frame)
            if title:
                continue
        if (
            kind == PP_PLACEHOLDER.SUBTITLE
            and getattr(shape, "has_text_frame", False)
            and not subtitle
        ):
            subtitle = _frame_plain(shape.text_frame)
            if subtitle:
                continue
        block = _shape_markdown(shape, width)
        if block:
            blocks.append(block)
    label = f"Diapositiva {number}"
    if hidden:
        label += " (oculta)"
    heading = f"## {label}. {_heading_text(title)}" if title else f"## {label}"
    parts = [heading, ""]
    if subtitle:
        parts.append(_paragraph_text(subtitle))
        parts.append("")
    for block in blocks:
        parts.append(block)
        parts.append("")
    notes = _notes_markdown(slide)
    if notes:
        parts.append(notes)
        parts.append("")
    return "\n".join(parts).rstrip() + "\n"


def pptx_to_markdown(pptx_path: Path, cell_width: int = DEFAULT_CELL_WIDTH) -> str:
    """Convierte cada diapositiva en una sección Markdown.

    Las tablas usan el mismo escape que Excel. Las imágenes sin texto
    alternativo y los diagramas SmartArt no se transcriben.
    """
    if cell_width < 0:
        cell_width = 0
    try:
        presentation = Presentation(str(pptx_path))
    except (PackageNotFoundError, BadZipFile, ValueError) as exc:
        raise RuntimeError(
            "No se pudo leer el PowerPoint. Si es .ppt antiguo, ábrelo y "
            "guárdalo como .pptx. Si tiene contraseña, quítala. "
            f"Detalle: {exc}"
        ) from exc

    total = len(presentation.slides)
    sections: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        hidden = _slide_is_hidden(presentation, index - 1)
        sections.append(_render_slide(slide, index, hidden, cell_width))
        console.print(f"  · {index}/{total}", markup=False)

    if sections:
        body = "\n\n".join(section.rstrip() for section in sections).strip() + "\n"
    else:
        body = "_La presentación no tiene diapositivas._\n"

    frontmatter = (
        "---\n"
        f"fuente: {_yaml_quote(pptx_path.name)}\n"
        f"diapositivas: {total}\n"
        'conversion: "powerpoint→markdown"\n'
        "generado_por: pdf2md\n"
        "---\n\n"
        f"# {pptx_path.stem}\n\n"
    )
    return frontmatter + body


# --------------------------------------------------------------------------- #
# Construcción del Markdown (PDF)
# --------------------------------------------------------------------------- #
def build_markdown(
    pdf_path: Path,
    ocr_mode: str,
    model: str,
    host: str,
    scanned: list[bool],
) -> str:
    """Genera el Markdown completo del PDF.

    Las páginas con texto se extraen con pymupdf4llm (sin Tesseract). Las páginas
    escaneadas (o todas, si ocr_mode='always') se transcriben con Ollama, una a
    una: no hay tope de páginas (el límite es el tiempo de OCR).
    """
    doc = pymupdf.open(pdf_path)
    total = doc.page_count
    parts: list[str] = []

    ocr_available = True
    if ocr_mode != "never":
        ocr_available = check_ollama(host, model)
        if not ocr_available:
            msg = f"Ollama/modelo '{model}' no disponible en {host}."
            if ocr_mode == "always" or any(scanned):
                console.print(
                    f"[yellow]{msg} Las páginas escaneadas quedarán vacías "
                    f"(¿está corriendo 'ollama serve'?).[/yellow]"
                )

    try:
        need_text_extract = ocr_mode != "always"
        chunks: list[dict] = []
        if need_text_extract:
            # use_ocr=False -> pymupdf4llm NO llama a Tesseract; solo extrae el
            # texto existente. Las páginas escaneadas se cubren con Ollama.
            chunks = pymupdf4llm.to_markdown(
                doc, page_chunks=True, use_ocr=False, show_progress=False
            )
        with TranscriptionDisplay(total, "Procesando") as view:
            for i in range(total):
                do_ocr = ocr_mode == "always" or (ocr_mode == "auto" and scanned[i])
                if do_ocr and ocr_available:
                    view.start_unit(f"OCR (visión) pág. {i + 1}/{total}")
                    try:
                        png = render_page_png(doc.load_page(i))
                        text = ocr_page_with_ollama(
                            png, model, host, on_chunk=view.feed
                        )
                    except requests.RequestException as exc:
                        console.print(f"[red]Error de OCR pág. {i + 1}: {exc}[/red]")
                        text = ""
                else:
                    view.start_unit(f"Texto pág. {i + 1}/{total}")
                    text = (
                        chunks[i].get("text", "") if i < len(chunks) else ""
                    ).strip()
                    view.show_text(text)
                parts.append(f"<!-- página {i + 1} -->\n\n{text}".rstrip())
                view.advance()
    finally:
        doc.close()

    body = "\n\n".join(parts).strip()
    frontmatter = (
        "---\n"
        f"fuente: {pdf_path.name}\n"
        f"paginas: {total}\n"
        f"ocr: {ocr_mode} (ollama:{model})\n"
        "generado_por: pdf2md\n"
        "---\n\n"
        f"# {pdf_path.stem}\n\n"
    )
    return frontmatter + body + "\n"


# --------------------------------------------------------------------------- #
# Construcción del Markdown (imágenes)
# --------------------------------------------------------------------------- #
def default_images_stem(images: list[Path]) -> str:
    if len(images) == 1:
        return images[0].stem
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return f"ocr-imagenes-{stamp}"


def build_images_markdown(
    images: list[Path],
    model: str,
    host: str,
) -> str:
    """OCR de cada imagen con Ollama y ensambla un único Markdown."""
    if not check_ollama(host, model):
        raise RuntimeError(
            f"Ollama/modelo '{model}' no disponible en {host}. "
            "¿Está corriendo 'ollama serve' y el modelo descargado?"
        )

    total = len(images)
    parts: list[str] = []

    with TranscriptionDisplay(total, "OCR imágenes") as view:
        for i, img_path in enumerate(images, start=1):
            view.start_unit(f"OCR {i}/{total}: {img_path.name}")
            try:
                image_bytes = img_path.read_bytes()
                text = ocr_image_bytes_with_ollama(
                    image_bytes, model, host, on_chunk=view.feed
                )
            except requests.RequestException as exc:
                console.print(f"[red]Error de OCR ({img_path.name}): {exc}[/red]")
                text = ""
            except OSError as exc:
                console.print(f"[red]No se pudo leer {img_path.name}: {exc}[/red]")
                text = ""
            parts.append(
                f"<!-- imagen {i}: {img_path.name} -->\n\n"
                f"## {img_path.name}\n\n"
                f"{text}".rstrip()
            )
            view.advance()

    body = "\n\n".join(parts).strip()
    fuentes_yaml = "\n".join(f"  - {p.name}" for p in images)
    title = images[0].stem if len(images) == 1 else "Transcripción OCR de imágenes"
    frontmatter = (
        "---\n"
        "fuentes:\n"
        f"{fuentes_yaml}\n"
        f"imagenes: {total}\n"
        f"ocr: ollama:{model}\n"
        "generado_por: pdf2md\n"
        "---\n\n"
        f"# {title}\n\n"
    )
    return frontmatter + body + "\n"


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def normalize_argv(argv: list[str]) -> list[str]:
    """Inserta el subcomando si el usuario pasó solo una ruta o flags."""
    if not argv:
        return argv
    if argv[0] in (
        "pdf", "images", "img", "imgs", "word", "docx",
        "excel", "xlsx", "xlsm", "pptx", "ppt", "slides",
        "md", "markdown", "-h", "--help",
    ):
        if argv[0] in ("img", "imgs"):
            return ["images", *argv[1:]]
        if argv[0] == "docx":
            return ["word", *argv[1:]]
        if argv[0] in ("xlsx", "xlsm"):
            return ["excel", *argv[1:]]
        if argv[0] in ("ppt", "slides"):
            return ["pptx", *argv[1:]]
        if argv[0] == "markdown":
            return ["md", *argv[1:]]
        return argv

    # Flags globales / de pdf sin subcomando → modo pdf (compatibilidad).
    if argv[0].startswith("-"):
        return ["pdf", *argv]

    first = Path(clean_dropped_path(argv[0]))
    if is_image_path(first):
        return ["images", *argv]
    if first.suffix.lower() == ".docx":
        return ["word", *argv]
    if is_excel_path(first) or first.suffix.lower() == ".xls":
        return ["excel", *argv]
    if is_pptx_path(first) or first.suffix.lower() == ".ppt":
        return ["pptx", *argv]
    if is_md_path(first):
        return ["md", *argv]
    return ["pdf", *argv]


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Transcribe un PDF o imágenes a Markdown o Word "
            "(OCR con Ollama / visión), convierte Word, Excel o PowerPoint "
            "→ Markdown, o un Markdown → Word. "
            "Sin límite de páginas."
        ),
    )
    sub = parser.add_subparsers(dest="command")

    pdf_p = sub.add_parser("pdf", help="Transcribir un PDF a Markdown o Word.")
    pdf_p.add_argument("pdf", nargs="?", help="Ruta del PDF (si se omite, se pregunta).")
    pdf_p.add_argument(
        "-o", "--output",
        help=(
            "Ruta o carpeta de salida. Por defecto: "
            f"{default_output_dir()} ({default_output_format()})"
        ),
    )
    pdf_p.add_argument(
        "--format", choices=["md", "docx"], default=None,
        dest="out_format",
        help=(
            "Formato de salida. Windows: docx. macOS/Linux: md. "
            "También se infiere por la extensión de -o."
        ),
    )
    pdf_p.add_argument(
        "--ocr", choices=["auto", "always", "never"], default="auto",
        help="Cuándo aplicar OCR con Ollama. 'auto' = solo páginas escaneadas.",
    )
    pdf_p.add_argument(
        "--model", default=None,
        help=(
            "Modelo de visión en Ollama. Si se omite, se usa PDF2MD_VISION_MODEL "
            "o el mejor modelo instalado según la RAM (9b en 16 GB)."
        ),
    )
    pdf_p.add_argument(
        "--host", default=DEFAULT_OLLAMA_HOST,
        help=f"Host de Ollama. Por defecto: {DEFAULT_OLLAMA_HOST}",
    )

    img_p = sub.add_parser(
        "images",
        aliases=["img", "imgs"],
        help="OCR de una o más imágenes → Markdown o Word.",
    )
    img_p.add_argument(
        "images",
        nargs="*",
        help="Rutas de imágenes (también puedes arrastrarlas al prompt).",
    )
    img_p.add_argument(
        "-o", "--output",
        help=(
            "Ruta o carpeta de salida. Por defecto: "
            f"{default_output_dir()} ({default_output_format()})"
        ),
    )
    img_p.add_argument(
        "--format", choices=["md", "docx"], default=None,
        dest="out_format",
        help=(
            "Formato de salida. Windows: docx. macOS/Linux: md. "
            "También se infiere por la extensión de -o."
        ),
    )
    img_p.add_argument(
        "--model", default=None,
        help=(
            "Modelo de visión en Ollama. Si se omite, se usa PDF2MD_VISION_MODEL "
            "o el mejor modelo instalado según la RAM (9b en 16 GB)."
        ),
    )
    img_p.add_argument(
        "--host", default=DEFAULT_OLLAMA_HOST,
        help=f"Host de Ollama. Por defecto: {DEFAULT_OLLAMA_HOST}",
    )

    word_p = sub.add_parser(
        "word",
        aliases=["docx"],
        help="Convertir un .docx a Markdown (requiere pandoc).",
    )
    word_p.add_argument(
        "docx",
        nargs="?",
        help="Ruta del .docx (si se omite, se pregunta).",
    )
    word_p.add_argument(
        "-o", "--output",
        help=f"Ruta o carpeta de salida (.md). Por defecto: {default_output_dir()}",
    )

    excel_p = sub.add_parser(
        "excel",
        aliases=["xlsx"],
        help="Convertir un .xlsx (o .xlsm) a tablas Markdown.",
    )
    excel_p.add_argument(
        "xlsx",
        nargs="?",
        help="Ruta del .xlsx (si se omite, se pregunta).",
    )
    excel_p.add_argument(
        "-o", "--output",
        help=f"Ruta o carpeta de salida (.md). Por defecto: {default_output_dir()}",
    )
    excel_p.add_argument(
        "--cell-width",
        type=int,
        default=DEFAULT_CELL_WIDTH,
        metavar="N",
        help=(
            "Máximo de caracteres por línea dentro de una celda. "
            f"Lo que sobre se parte con <br> (predet. {DEFAULT_CELL_WIDTH}). "
            "0 = no partir."
        ),
    )

    pptx_p = sub.add_parser(
        "pptx",
        aliases=["slides"],
        help="Convertir un .pptx (o .pptm) a Markdown.",
    )
    pptx_p.add_argument(
        "pptx",
        nargs="?",
        help="Ruta del .pptx (si se omite, se pregunta).",
    )
    pptx_p.add_argument(
        "-o", "--output",
        help=f"Ruta o carpeta de salida (.md). Por defecto: {default_output_dir()}",
    )
    pptx_p.add_argument(
        "--cell-width",
        type=int,
        default=DEFAULT_CELL_WIDTH,
        metavar="N",
        help=(
            "Máximo de caracteres por línea dentro de una celda de tabla. "
            f"Lo que sobre se parte con <br> (predet. {DEFAULT_CELL_WIDTH}). "
            "0 = no partir."
        ),
    )

    md_p = sub.add_parser(
        "md",
        aliases=["markdown"],
        help="Convertir un .md a Word (.docx).",
    )
    md_p.add_argument(
        "md",
        nargs="?",
        help="Ruta del .md (si se omite, se pregunta).",
    )
    md_p.add_argument(
        "-o", "--output",
        help=(
            "Ruta o carpeta de salida (.docx). "
            f"Por defecto: {default_output_dir()}"
        ),
    )

    return parser.parse_args(argv)


def run_pdf(args: argparse.Namespace) -> int:
    pdf_path = Path(clean_dropped_path(args.pdf)) if args.pdf else prompt_for_pdf()
    if not pdf_path.exists() or pdf_path.suffix.lower() != ".pdf":
        console.print(f"[red]Archivo PDF inválido:[/red] {pdf_path}")
        return 1

    fmt = resolve_output_format(getattr(args, "out_format", None), args.output)
    out_path = resolve_output_path(pdf_path.stem, args.output, fmt)

    try:
        scanned = detect_scanned_pages(pdf_path)
    except Exception as exc:  # noqa: BLE001
        console.print(f"[red]No se pudo abrir el PDF:[/red] {exc}")
        return 1
    n_scanned = sum(scanned)

    console.print(f"[cyan]Entrada :[/cyan] {pdf_path}")
    console.print(f"[cyan]Salida  :[/cyan] {out_path}")
    console.print(
        f"[cyan]Páginas :[/cyan] {len(scanned)} "
        f"({n_scanned} escaneadas -> OCR; sin tope de n)"
    )
    console.print(f"[cyan]OCR     :[/cyan] {args.ocr} · ollama · {args.model}")
    console.print(f"[cyan]Formato :[/cyan] {fmt}")

    try:
        markdown = build_markdown(pdf_path, args.ocr, args.model, args.host, scanned)
    except Exception as exc:  # noqa: BLE001
        console.print(f"[red]Error al procesar el PDF:[/red] {exc}")
        return 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    write_output(markdown, out_path, fmt)

    kind = "Word" if fmt == "docx" else "Markdown"
    console.print(
        Panel.fit(
            f"[green]Listo[/green] -> {out_path}\n"
            f"{len(markdown):,} caracteres de transcripción ({kind}).",
            border_style="green",
        )
    )
    return 0


def run_images(args: argparse.Namespace) -> int:
    if args.images:
        # Permite tanto args separados como una sola línea con varias rutas.
        raw_joined = " ".join(args.images)
        # Si cada arg ya es una ruta limpia de argparse, preferimos eso.
        candidates = [Path(clean_dropped_path(p)) for p in args.images]
        # Si el usuario pegó todo en un solo token con escapes, re-parsear.
        if len(candidates) == 1 and ("\\" in args.images[0] or " " in args.images[0]):
            candidates = parse_dropped_paths(raw_joined)
        images, errors = validate_image_paths(candidates)
        for msg in errors:
            console.print(f"[red]{msg}[/red]")
        if not images:
            console.print("[red]No hay imágenes válidas para procesar.[/red]")
            return 1
    else:
        images = prompt_for_images()

    fmt = resolve_output_format(getattr(args, "out_format", None), args.output)
    out_path = resolve_output_path(default_images_stem(images), args.output, fmt)

    console.print(f"[cyan]Imágenes:[/cyan] {len(images)}")
    for img in images:
        console.print(f"  · {img}")
    console.print(f"[cyan]Salida  :[/cyan] {out_path}")
    console.print(f"[cyan]OCR     :[/cyan] ollama · {args.model}")
    console.print(f"[cyan]Formato :[/cyan] {fmt}")

    try:
        markdown = build_images_markdown(images, args.model, args.host)
    except Exception as exc:  # noqa: BLE001
        console.print(f"[red]Error al procesar las imágenes:[/red] {exc}")
        return 1

    write_output(markdown, out_path, fmt)

    kind = "Word" if fmt == "docx" else "Markdown"
    console.print(
        Panel.fit(
            f"[green]Listo[/green] -> {out_path}\n"
            f"{len(images)} imagen(es) · {len(markdown):,} caracteres ({kind}).",
            border_style="green",
        )
    )
    return 0


def run_word(args: argparse.Namespace) -> int:
    docx_path = Path(clean_dropped_path(args.docx)) if args.docx else prompt_for_docx()
    if not docx_path.exists() or docx_path.suffix.lower() != ".docx":
        console.print(f"[red]Archivo Word inválido:[/red] {docx_path}")
        return 1

    out_path = resolve_output_path(docx_path.stem, args.output, "md")

    console.print(f"[cyan]Entrada :[/cyan] {docx_path}")
    console.print(f"[cyan]Salida  :[/cyan] {out_path}")
    console.print("[cyan]Motor   :[/cyan] pandoc (Word → Markdown)")

    try:
        markdown = docx_to_markdown(docx_path)
    except Exception as exc:  # noqa: BLE001
        console.print(f"[red]Error al convertir el Word:[/red] {exc}")
        return 1

    write_output(markdown, out_path, "md")

    console.print(
        Panel.fit(
            f"[green]Listo[/green] -> {out_path}\n"
            f"{len(markdown):,} caracteres de Markdown.",
            border_style="green",
        )
    )
    return 0


def run_excel(args: argparse.Namespace) -> int:
    if args.xlsx:
        xlsx_path = Path(clean_dropped_path(args.xlsx))
    else:
        xlsx_path = prompt_for_xlsx()
    if not xlsx_path.exists():
        console.print(f"[red]No existe:[/red] {xlsx_path}")
        return 1
    if xlsx_path.suffix.lower() == ".xls":
        console.print(
            "[red]El formato .xls (Excel 97-2003) no está soportado.[/red] "
            "Ábrelo y guárdalo como .xlsx."
        )
        return 1
    if not is_excel_path(xlsx_path):
        console.print(f"[red]Archivo Excel inválido:[/red] {xlsx_path}")
        return 1

    cell_width = args.cell_width if args.cell_width > 0 else 0
    out_path = resolve_output_path(xlsx_path.stem, args.output, "md")

    console.print(f"[cyan]Entrada :[/cyan] {xlsx_path}")
    console.print(f"[cyan]Salida  :[/cyan] {out_path}")
    ancho = "sin corte" if cell_width == 0 else str(cell_width)
    console.print(f"[cyan]Celdas  :[/cyan] ancho {ancho}")

    try:
        markdown = xlsx_to_markdown(xlsx_path, cell_width=cell_width)
    except Exception as exc:  # noqa: BLE001
        console.print(f"[red]Error al convertir el Excel:[/red] {exc}")
        return 1

    write_output(markdown, out_path, "md")

    console.print(
        Panel.fit(
            f"[green]Listo[/green] -> {out_path}\n"
            f"{len(markdown):,} caracteres de Markdown.",
            border_style="green",
        )
    )
    return 0


def run_pptx(args: argparse.Namespace) -> int:
    if args.pptx:
        pptx_path = Path(clean_dropped_path(args.pptx))
    else:
        pptx_path = prompt_for_pptx()
    if not pptx_path.exists():
        console.print(f"[red]No existe:[/red] {pptx_path}")
        return 1
    if pptx_path.suffix.lower() == ".ppt":
        console.print(
            "[red]El formato .ppt (PowerPoint 97-2003) no está soportado.[/red] "
            "Ábrelo y guárdalo como .pptx."
        )
        return 1
    if not is_pptx_path(pptx_path):
        console.print(f"[red]Archivo PowerPoint inválido:[/red] {pptx_path}")
        return 1

    cell_width = args.cell_width if args.cell_width > 0 else 0
    out_path = resolve_output_path(pptx_path.stem, args.output, "md")

    console.print(f"[cyan]Entrada :[/cyan] {pptx_path}")
    console.print(f"[cyan]Salida  :[/cyan] {out_path}")
    ancho = "sin corte" if cell_width == 0 else str(cell_width)
    console.print(f"[cyan]Tablas  :[/cyan] ancho {ancho}")

    try:
        markdown = pptx_to_markdown(pptx_path, cell_width=cell_width)
    except Exception as exc:  # noqa: BLE001
        console.print(f"[red]Error al convertir el PowerPoint:[/red] {exc}")
        return 1

    write_output(markdown, out_path, "md")

    console.print(
        Panel.fit(
            f"[green]Listo[/green] -> {out_path}\n"
            f"{len(markdown):,} caracteres de Markdown.",
            border_style="green",
        )
    )
    return 0


def run_md(args: argparse.Namespace) -> int:
    md_path = Path(clean_dropped_path(args.md)) if args.md else prompt_for_md()
    if not md_path.exists():
        console.print(f"[red]No existe:[/red] {md_path}")
        return 1
    if not is_md_path(md_path):
        console.print(f"[red]Archivo Markdown inválido:[/red] {md_path}")
        return 1

    out_path = resolve_output_path(md_path.stem, args.output, "docx")
    if out_path.suffix.lower() != ".docx":
        out_path = out_path.with_suffix(".docx")

    console.print(f"[cyan]Entrada :[/cyan] {md_path}")
    console.print(f"[cyan]Salida  :[/cyan] {out_path}")
    console.print("[cyan]Motor   :[/cyan] python-docx (Markdown → Word)")

    try:
        markdown = md_path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        console.print("[red]El Markdown no está en UTF-8.[/red]")
        return 1
    except OSError as exc:
        console.print(f"[red]No se pudo leer el Markdown:[/red] {exc}")
        return 1

    try:
        markdown_to_docx(markdown, out_path, join_lines=True)
    except Exception as exc:  # noqa: BLE001
        console.print(f"[red]Error al convertir el Markdown:[/red] {exc}")
        return 1

    console.print(
        Panel.fit(
            f"[green]Listo[/green] -> {out_path}\n"
            f"{len(markdown):,} caracteres de Markdown → Word.",
            border_style="green",
        )
    )
    return 0


def main(argv: list[str] | None = None) -> int:
    try:
        return _main(argv)
    except KeyboardInterrupt:
        console.print("\n[yellow]Cancelado.[/yellow]")
        return 130


def _main(argv: list[str] | None = None) -> int:
    raw = list(argv if argv is not None else sys.argv[1:])
    interactive = not raw
    if not raw:
        mode = prompt_for_mode()
        if mode is None:
            console.print("[yellow]Cancelado.[/yellow]")
            return 0
        raw = [mode]

    args = parse_args(normalize_argv(raw))
    if args.command in ("xlsx", "xlsm"):
        args.command = "excel"
    if args.command in ("ppt", "slides"):
        args.command = "pptx"
    if args.command == "markdown":
        args.command = "md"
    if args.command == "docx":
        args.command = "word"
    if args.command in ("img", "imgs"):
        args.command = "images"
    if not args.command:
        mode = prompt_for_mode()
        if mode is None:
            console.print("[yellow]Cancelado.[/yellow]")
            return 0
        args = parse_args([mode])
        interactive = True

    if args.command == "word":
        return run_word(args)
    if args.command == "excel":
        return run_excel(args)
    if args.command == "pptx":
        return run_pptx(args)
    if args.command == "md":
        return run_md(args)

    if interactive and not getattr(args, "out_format", None) and os.name != "nt":
        args.out_format = Prompt.ask(
            "[cyan]Formato[/cyan]",
            choices=["md", "docx"],
            default="md",
        )

    args.model = resolve_vision_model(args.host, args.model)
    warn_if_model_too_heavy(args.model)

    if args.command == "pdf":
        return run_pdf(args)
    if args.command == "images":
        return run_images(args)

    console.print(
        "[red]Comando no reconocido. "
        "Usa 'pdf', 'images', 'word', 'excel', 'pptx' o 'md'.[/red]"
    )
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
